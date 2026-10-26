"""Read-only XLSX ingestion. Formula caches are evidence, never recalculation."""
from __future__ import annotations

import argparse
import os
import hashlib
import json
import math
import shutil
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
_local_config = ROOT / "config.local.json"
_local = json.loads(_local_config.read_text(encoding="utf-8")) if _local_config.exists() else {}
SOURCE = Path(os.environ.get("SIGMA_SOURCE_XLSX") or _local.get("source_xlsx") or ROOT / "data" / "source.xlsx")
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
ERRORS = {"#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"}


def scalar(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def number(value):
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) else None


def key(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip() if value is not None else None


def inspect_formula_caches(path: Path, sheet_names: list[str]):
    summaries = []
    formula_cells = {}
    with ZipFile(path) as z:
        # Resolve worksheet relationships rather than assuming sheet order is file order.
        relns = "http://schemas.openxmlformats.org/package/2006/relationships"
        rels = {e.attrib["Id"]: e.attrib["Target"] for e in ET.fromstring(z.read("xl/_rels/workbook.xml.rels")).findall(f"{{{relns}}}Relationship")}
        workbook = ET.fromstring(z.read("xl/workbook.xml"))
        for sheet in workbook.findall("m:sheets/m:sheet", NS):
            name = sheet.attrib["name"]
            rid = sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
            target = rels[rid]
            entry = target.lstrip("/") if target.startswith("/") else "xl/" + target
            root = ET.fromstring(z.read(entry))
            formulas, missing, errors = 0, 0, Counter()
            samples = []
            formula_cells[name] = set()
            for c in root.findall(".//m:sheetData/m:row/m:c", NS):
                f, v = c.find("m:f", NS), c.find("m:v", NS)
                if f is not None:
                    formulas += 1
                    formula_cells[name].add(c.attrib["r"])
                    if v is None or v.text is None:
                        missing += 1
                if c.attrib.get("t") == "e":
                    err = v.text if v is not None else "unknown"
                    errors[err] += 1
                    if len(samples) < 6:
                        samples.append({"cell": c.attrib["r"], "error": err})
            summaries.append({"sheet": name, "formula_count": formulas, "empty_formula_cache": missing, "cached_errors": dict(errors), "error_examples": samples})
    return summaries, formula_cells


def ingest(path: Path = SOURCE, root: Path = ROOT):
    path = Path(path).resolve(strict=True)
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    raw = root / "data" / "raw" / f"{sha[:16]}.xlsx"
    raw.parent.mkdir(parents=True, exist_ok=True)
    if not raw.exists():
        shutil.copy2(path, raw)
    if hashlib.sha256(raw.read_bytes()).hexdigest() != sha:
        raise ValueError("La copia de origen no coincide con la huella del archivo.")
    wb = openpyxl.load_workbook(raw, read_only=True, data_only=True, keep_links=False)
    formulas, formula_cells = inspect_formula_caches(raw, wb.sheetnames)
    required = {"BD", "BD Comercial Occidente", "Modelos y Marcas", "Capacidades Electricas", "I. Solicitud de Información"}
    if required - set(wb.sheetnames):
        raise ValueError("Faltan hojas requeridas: " + ", ".join(sorted(required - set(wb.sheetnames))))

    master = defaultdict(list)
    master_rows = 0
    for ix, row in enumerate(wb["BD"].iter_rows(min_row=3, max_col=40, values_only=True), 3):
        if row[0] is None:
            continue
        master_rows += 1
        master[key(row[0])].append({"row": ix, "year": scalar(row[5]), "manufacturer": row[16], "capacity_raw": scalar(row[26]), "status": row[29], "branch": row[8]})

    sheet = wb["BD Comercial Occidente"]
    headers = [scalar(c.value) for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    fields = {"region": 0, "branch": 1, "equipment": 2, "capacity_kg": 3, "operation": 4, "description": 5, "fuel": 6, "route": 7, "route_sap": 8, "route_type": 9, "reported_daily_load_kg": 10, "reported_daily_km": 11, "reported_max_km": 12, "reported_weekly_km": 13, "brand": 16, "model": 17}
    vehicles = []
    daily = []
    for ix, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if row[2] is None:
            continue
        item = {f: scalar(row[j]) for f, j in fields.items()}
        item["equipment"] = key(row[2])
        item["record_id"] = f"commercial:{ix}"
        item["source_row"] = ix
        item["source_sheet"] = sheet.title
        item["source_sha256"] = sha
        item["master_matches"] = master.get(item["equipment"], [])
        item["source_cells"] = {f: f"{openpyxl.utils.get_column_letter(j+1)}{ix}" for f, j in fields.items()}
        item["formula_fields"] = [f for f, cell in item["source_cells"].items() if cell in formula_cells[sheet.title]]
        item["issues"] = []
        item["raw_values"] = {f: scalar(row[j]) for f, j in fields.items()}
        for f, value in list(item.items()):
            if isinstance(value, str) and value in ERRORS:
                item["issues"].append(f"{f}: {value}")
                item[f] = None
        for f in ["capacity_kg", "reported_daily_load_kg", "reported_daily_km", "reported_max_km", "reported_weekly_km"]:
            value = item[f]
            item[f] = number(value)
            if value is not None and item[f] is None:
                item["issues"].append(f"{f}: valor no numérico {value!r}; conservado en raw_values")
        for f in ["capacity_kg", "reported_daily_km", "reported_daily_load_kg"]:
            if item[f] is None:
                item["issues"].append(f"{f}: pendiente")
        if len(item["master_matches"]) != 1:
            item["issues"].append(f"BD: {len(item['master_matches'])} coincidencias")
        vehicles.append(item)
        for j, header in enumerate(headers):
            if isinstance(header, str) and header.startswith("2026-") and row[j] is not None:
                daily.append({"record_id": item["record_id"], "equipment": item["equipment"], "date": header[:10], "value_raw": scalar(row[j]), "unit": "km por confirmar", "source_cell": f"{openpyxl.utils.get_column_letter(j+1)}{ix}"})

    sites = []
    for ix, row in enumerate(wb["Capacidades Electricas"].iter_rows(min_row=2, max_col=4, values_only=True), 2):
        if row[0] is not None:
            sites.append({"record_id": f"site:{ix}", "name": str(row[0]).strip(), "connected_kw": number(row[1]), "contracted_kw": number(row[2]), "historical_peak_kw": number(row[3]), "available_kw": None, "source_row": ix, "source_sheet": "Capacidades Electricas", "validation": "Capacidad adicional sin verificar"})
    models = []
    for ix, row in enumerate(wb["Modelos y Marcas"].iter_rows(min_row=3, max_col=7, values_only=True), 3):
        if row[3] is not None:
            models.append({"brand": row[3], "model": row[4], "category_raw": row[5], "category_header": "Categoria Sigma (ton)", "capacity_kg": None, "powertrain": row[6], "source_row": ix, "status": "Unidades y especificaciones por verificar"})
    notes = []
    for ix, row in enumerate(wb["I. Solicitud de Información"].iter_rows(min_row=13, max_row=20, max_col=10, values_only=True), 13):
        notes.append({"item": row[0], "request": row[4], "sigma_note": row[9], "source": f"I. Solicitud de Información!J{ix}"})
    duplicates = {k: n for k, n in Counter(v["equipment"] for v in vehicles).items() if n > 1}
    site_duplicates = {k: n for k, n in Counter(s["name"].casefold() for s in sites).items() if n > 1}
    issues = [
        {"id": "regions", "level": "review", "title": "Confirmar alcance regional", "detail": "El alcance occidente/noroeste debe relacionarse con Bajío, Pacífico y Noroeste; no hay filtro implícito.", "source": "BD Comercial Occidente!A2:A966"},
        {"id": "units", "level": "blocking", "title": "Unidades del catálogo sin confirmar", "detail": "El encabezado indica toneladas, con valores 600–5000. No se convierten ni se usan como capacidad útil.", "source": "Modelos y Marcas!F2:F41"},
        {"id": "cycle", "level": "blocking", "title": "Falta el ciclo de conducción", "detail": "Los indicadores y registros de kilometraje no incluyen velocidad temporal, pendiente ni aperturas. No se reconstruye una ruta real a partir de promedios.", "source": "BD Comercial Occidente!H1:N1"},
        {"id": "thermal", "level": "blocking", "title": "Faltan caja y equipo frigorífico", "detail": "Temperatura solicitada disponible; aislamiento, geometría, mapas, producto y aperturas pendientes.", "source": "I. Solicitud de Información!J14"},
        {"id": "diesel", "level": "blocking", "title": "Consumo diésel pendiente", "detail": "No se calcula ahorro frente a diésel sin una línea base explícita y comparable.", "source": "I. Solicitud de Información!J16"},
        {"id": "site", "level": "blocking", "title": "Capacidad adicional del sitio sin verificar", "detail": "La demanda contratada y el pico histórico no acreditan capacidad técnica ni coincidencia horaria.", "source": "Capacidades Electricas!A1:D113"},
        {"id": "cache", "level": "review", "title": "Resultados de fórmulas almacenados", "detail": "Lectura de valores guardados por Excel; no recalculados. Se conservan errores y procedencia.", "source": "Libro completo"},
    ]
    for name, count, detail in [
        ("duplicate_equipment", len(duplicates), "Identificadores repetidos en la base comercial; conservar filas y resolver antes de unir."),
        ("duplicate_sites", len(site_duplicates), "Nombres de sitios repetidos; no agregar ni elegir una fila automáticamente."),
        ("missing_master", sum(not v["master_matches"] for v in vehicles), "Registros comerciales sin coincidencia en BD."),
        ("error_model", sum(any("model: #" in x for x in v["issues"]) for v in vehicles), "Modelos no resueltos por las fórmulas del libro."),
    ]:
        if count:
            issues.append({"id": name, "level": "review", "title": detail, "count": count, "detail": detail, "source": "Bases del libro"})
    catalogue = {"source": {"name": path.name, "path": str(path), "sha256": sha, "snapshot": str(raw.relative_to(root)), "bytes": path.stat().st_size, "imported_at": datetime.now(timezone.utc).isoformat(), "cache_policy": "Valores almacenados; no recalculados"}, "sheets": [{"name": s.title, "max_row": s.max_row, "max_column": s.max_column} for s in wb], "vehicles": vehicles, "sites": sites, "models": models, "notes": notes, "issues": issues, "formula_audit": formulas, "summary": {"commercial_records": len(vehicles), "unique_equipment": len(set(v["equipment"] for v in vehicles)), "master_records": master_rows, "sites_records": len(sites), "models_records": len(models), "daily_observations": len(daily), "regions": dict(Counter(v["region"] for v in vehicles)), "operation": dict(Counter(v["operation"] for v in vehicles)), "duplicate_equipment": duplicates, "duplicate_sites": site_duplicates, "missing_capacity": sum(v["capacity_kg"] is None for v in vehicles), "missing_daily_km": sum(v["reported_daily_km"] is None for v in vehicles)}}
    wb.close()
    processed = root / "data" / "processed"
    processed.mkdir(parents=True, exist_ok=True)
    (processed / "catalogue.json").write_text(json.dumps(catalogue, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    (processed / "daily_raw.json").write_text(json.dumps(daily, ensure_ascii=False, allow_nan=False), encoding="utf-8")
    dictionary = [{"field": f, "header": headers[j], "sheet": sheet.title, "column": openpyxl.utils.get_column_letter(j+1), "meaning_status": "Confirmar periodo y semántica" if f.startswith("reported_") else "Proporcionado", "transform": "Errores a null con registro; unidades originales preservadas"} for f, j in fields.items()]
    (processed / "dictionary.json").write_text(json.dumps(dictionary, ensure_ascii=False, indent=2), encoding="utf-8")
    if hashlib.sha256(path.read_bytes()).hexdigest() != sha:
        raise ValueError("El archivo original cambió durante la importación; repetir lectura.")
    return catalogue


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=SOURCE)
    args = parser.parse_args()
    result = ingest(args.source)
    print(json.dumps({"sha256": result["source"]["sha256"], "summary": result["summary"], "formula_audit": result["formula_audit"]}, ensure_ascii=False, indent=2))
